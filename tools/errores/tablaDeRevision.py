import json, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open(sys.argv[1]))
filas, errores = d["filas"], d["errores"]
F = "Arial"
bold = Font(name=F, bold=True, color="FFFFFF")
normal = Font(name=F, size=10)
cab_fill = PatternFill("solid", fgColor="7A4319")
amarillo = PatternFill("solid", fgColor="FFF4C2")
gris = PatternFill("solid", fgColor="F2EFE9")
choque_fill = PatternFill("solid", fgColor="FBE3D6")
fino = Side(style="thin", color="D9D2C5")
borde = Border(bottom=fino)
RECTA = {"representa_en_recta", "lee_la_recta", "representa_en_recta_graduada", "lee_la_recta_graduada"}
SIN = {
  "representa_en_recta": "La respuesta es un dibujo (puntos en la recta). Ningún error del catálogo predice dónde los pondría el alumno.",
  "representa_en_recta_graduada": "Igual que la anterior, con otra escala.",
  "lee_la_recta": "Leer marcas: el fallo típico (contar mal las marcas, leer al revés los negativos) no está en el catálogo.",
  "lee_la_recta_graduada": "Igual: el error de escala (cada marca vale 5, no 1) no está en el catálogo.",
  "series_numericas": "Cruzar el cero en una serie (1, -1, -3…) no corresponde a ningún error del catálogo tal como está escrito.",
  "potencias_base_entera": "El catálogo NO tiene errores de potencias: (-2)^3 = 8 o (-3)^2 = -9 no se pueden diagnosticar.",
  "pares_con_y_sin_parentesis": "Esta batería existe justo para el error -3^2 = 9 (el menos dentro de la potencia), y ese error NO está en el catálogo.",
}

wb = Workbook()
lee = wb.active
lee.title = "Léeme"
texto = [
  ("Respuestas-trampa — Enteros 1.º ESO", True),
  ("Qué escribiría un alumno con cada error típico, en cada apartado. Nunca se imprime: es lo que permitirá corregir diciendo POR QUÉ ha fallado.", False),
  ("", False),
  ("Cómo leerla", True),
  ("• Pestaña «Trampas»: un apartado por fila (2 ejemplos de cada batería, 4 apartados cada uno). Columnas E1…E13: la respuesta que da ese error. Vacío = ese apartado no detecta ese error (con el error sale bien, o no aplica).", False),
  ("• «Choques»: dos errores que dan la MISMA respuesta. Ese apartado no sirve para distinguirlos: el diagnóstico dirá «uno de estos dos».", False),
  ("• Pestaña «Por error»: cuántos apartados detectan cada error. Pestaña «Sin detección»: baterías donde no hay ninguna trampa y por qué.", False),
  ("", False),
  ("Qué tienes que hacer tú (celdas amarillas)", True),
  ("• En «Trampas», columna Revisión: OK / Mal / Dudoso. Si es «Mal», escribe en Comentario qué escribiría de verdad el alumno.", False),
  ("• Ejemplo: fila «-8 + 5 = ___», E2 = -13 → OK (es tu ejemplo del catálogo).", False),
  ("", False),
  ("Decisiones mías que tienes que confirmar", True),
  ("1. Error 7 (quitar paréntesis): uso tu ejemplo «8 - (-3 + 5) = 8 - 3 + 5», o sea, quita el paréntesis SIN cambiar ningún signo. La descripción dice también «o solo el primero», que da otra respuesta (8 + 3 + 5 = 16). ¿Cuál es la más frecuente en tus alumnos?", False),
  ("2. Errores 2 y 3 (signos distintos) solo cuando hay un negativo ESCRITO: en «50 - 14» no hay error de enteros posible.", False),
  ("3. Un error de concepto se aplica en TODA la expresión (es una creencia, no un despiste). Si solo falla una vez, es el 15 (descuido), que no se predice.", False),
  ("4. Error 4 al ordenar: el alumno sabe que los negativos van antes, pero los ordena por su número sin signo (-1 < -3 < -8).", False),
  ("5. Error 13 (el cero): ninguna batería pregunta si el 0 es positivo o negativo, así que no aparece.", False),
]
for i, (t, negrita) in enumerate(texto, 1):
  c = lee.cell(row=i, column=1, value=t)
  c.font = Font(name=F, bold=negrita, size=13 if i == 1 else 10)
  c.alignment = Alignment(wrap_text=True, vertical="top")
lee.column_dimensions["A"].width = 120

ws = wb.create_sheet("Trampas")
cols = ["Obj.", "Batería", "Apartado", "Correcta"] + [f"E{e['numero']}" for e in errores] + ["Detecta", "Choques", "Revisión", "Comentario"]
for j, h in enumerate(cols, 1):
  c = ws.cell(row=1, column=j, value=h)
  c.font = bold; c.fill = cab_fill; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for k, e in enumerate(errores):
  ws.cell(row=1, column=5 + k).comment = Comment(f"Error {e['numero']} ({e['categoria']}): {e['corto']}", "TutorDigital")
primeraE, ultimaE = 5, 4 + len(errores)
colDet, colCho, colRev, colCom = ultimaE + 1, ultimaE + 2, ultimaE + 3, ultimaE + 4
anterior = None
for i, f in enumerate(filas, 2):
  enun = "(recta dibujada)" if f["clave"] in RECTA else f["texto"]
  vals = [f["objetivo"], f["clave"], enun, f["correcta"]]
  for j, v in enumerate(vals, 1):
    ws.cell(row=i, column=j, value=v).font = normal
  for k, e in enumerate(errores):
    r = f["trampas"].get(str(e["numero"]))
    c = ws.cell(row=i, column=5 + k, value=r)
    c.font = normal; c.alignment = Alignment(horizontal="center")
  L1, L2 = get_column_letter(primeraE), get_column_letter(ultimaE)
  ws.cell(row=i, column=colDet, value=f"=COUNTA({L1}{i}:{L2}{i})").font = normal
  porResp = defaultdict(list)
  for e, r in f["trampas"].items(): porResp[r].append(int(e))
  choques = "; ".join("=".join(f"E{x}" for x in sorted(es)) + f" ({r})" for r, es in porResp.items() if len(es) > 1)
  c = ws.cell(row=i, column=colCho, value=choques or None); c.font = normal
  if choques: c.fill = choque_fill
  for col in (colRev, colCom):
    ws.cell(row=i, column=col).fill = amarillo
  if f["clave"] != anterior:
    for j in range(1, colCom + 1): ws.cell(row=i, column=j).border = Border(top=Side(style="thin", color="7A4319"))
  anterior = f["clave"]
dv = DataValidation(type="list", formula1='"OK,Mal,Dudoso"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"{get_column_letter(colRev)}2:{get_column_letter(colRev)}{len(filas) + 1}")
widths = {1: 5, 2: 26, 3: 44, 4: 16}
for j in range(1, colCom + 1):
  ws.column_dimensions[get_column_letter(j)].width = widths.get(j, 9)
ws.column_dimensions[get_column_letter(colCho)].width = 22
ws.column_dimensions[get_column_letter(colRev)].width = 11
ws.column_dimensions[get_column_letter(colCom)].width = 40
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(colCom)}{len(filas) + 1}"
ws.row_dimensions[1].height = 22

pe = wb.create_sheet("Por error")
hdr = ["Nº", "Error", "Categoría", "Apartados que lo detectan", "Baterías donde aparece"]
for j, h in enumerate(hdr, 1):
  c = pe.cell(row=1, column=j, value=h); c.font = bold; c.fill = cab_fill; c.alignment = Alignment(wrap_text=True)
n = len(filas) + 1
for k, e in enumerate(errores, 2):
  col = get_column_letter(5 + k - 2)
  baterias = sorted({f["clave"] for f in filas if str(e["numero"]) in f["trampas"]})
  vals = [e["numero"], e["corto"], e["categoria"], f"=COUNTA(Trampas!{col}2:{col}{n})", ", ".join(baterias) or "— ninguna —"]
  for j, v in enumerate(vals, 1):
    c = pe.cell(row=k, column=j, value=v); c.font = normal; c.alignment = Alignment(wrap_text=True, vertical="top")
  if not baterias:
    for j in range(1, 6): pe.cell(row=k, column=j).fill = choque_fill
tot = len(errores) + 2
pe.cell(row=tot, column=3, value="Apartados en la muestra").font = Font(name=F, bold=True)
pe.cell(row=tot, column=4, value=f"=COUNTA(Trampas!A2:A{n})").font = Font(name=F, bold=True)
for j, w in enumerate([5, 55, 15, 14, 60], 1): pe.column_dimensions[get_column_letter(j)].width = w

sd = wb.create_sheet("Sin detección")
for j, h in enumerate(["Batería", "Por qué no hay trampas"], 1):
  c = sd.cell(row=1, column=j, value=h); c.font = bold; c.fill = cab_fill
claves = []
for f in filas:
  if f["clave"] not in claves: claves.append(f["clave"])
r = 2
for cl in claves:
  if any(f["trampas"] for f in filas if f["clave"] == cl): continue
  sd.cell(row=r, column=1, value=cl).font = normal
  c = sd.cell(row=r, column=2, value=SIN.get(cl, "Ningún error del catálogo cambia la respuesta de estos apartados.")); c.font = normal; c.alignment = Alignment(wrap_text=True)
  r += 1
sd.column_dimensions["A"].width = 30; sd.column_dimensions["B"].width = 100
wb.save(sys.argv[2])
print("ok", len(filas))
